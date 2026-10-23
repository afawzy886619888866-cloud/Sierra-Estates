#!/usr/bin/env python3
"""Sierra Estates — Owner Verification Campaign worklist builder (repo edition).

Regenerates the Top-N owner-verification worklist from the in-repo Google Sheet
harvest (apps/sierra-estates-realty/data/online-sheet-units.json) and writes an
xlsx + csv OUTSIDE the repo tree (owner-PII must never be committed).

Scoring (0-100):
  contact      25   valid normalized phone (hard gate to enter the queue)
  price        25   plausible for deal type (rent 3k-500k/mo; sale 0.8M-300M)
  completeness 25   5 pts each: compound / type / beds / baths / area
  freshness    15   lastUpdate <=30d:15  <=90d:10  <=180d:5  else 0
  compound     10   tier-A compound (Mivida, Palm Hills, Rehab, ...)
Tiers: P1 >= 70, P2 55-69, P3 < 55.

Usage:
  python scripts/data-audit/build_verification_campaign.py [--top 500] \
      [--out-dir ../sierra-campaign]

Dependencies: openpyxl (pip install openpyxl). No other requirements.
"""
import argparse
import csv
import json
import os
import re
import sys
from collections import Counter
from datetime import datetime, timezone

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.formatting.rule import DataBarRule
    from openpyxl.chart import BarChart, Reference
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
SRC = os.path.join(REPO, "apps", "sierra-estates-realty", "data",
                   "online-sheet-units.json")

PRIMARY = "1B2A4A"
NEUTRAL_900 = "37352F"
NEUTRAL_600 = "8C8A84"
NEUTRAL_100 = "F7F7F5"

TIER_A_COMPOUNDS = {
    "mivida", "palm hills", "rehab", "madinaty", "hyde park", "mountain view",
    "katameya", "katameya dwar", "sodic", "fifth square", "the square",
    "lake view", "beverly", "azzar", "la vista", "telal", "grassridge",
    "grand heights", "dream land", "waterway", "karmell",
}
STATUS_OPTIONS = ("Pending,Call Again,Confirmed Available,Price Updated,"
                  "Sold / Rented,Wrong Number,Opted Out")
MONTHS = {"jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
          "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12}


def clean_phone(raw):
    if raw is None:
        return None
    digits = re.sub(r"\D", "", str(raw).split("@")[0].split(".")[0])
    if digits.startswith("201") and len(digits) == 12:
        return digits
    if digits.startswith("01") and len(digits) == 11:
        return "2" + digits
    if digits.startswith("1") and len(digits) == 10:
        return "20" + digits
    return digits if len(digits) >= 10 else None


def parse_date(raw):
    if not raw:
        return None
    m = re.match(r"^(\d{1,2})-([A-Za-z]{3})-(\d{2,4})$", str(raw).strip())
    if m:
        day = int(m.group(1))
        mon = MONTHS.get(m.group(2).lower())
        year = int(m.group(3))
        if mon:
            year += 2000 if year < 100 else 0
            try:
                return datetime(year, mon, day)
            except ValueError:
                return None
    return None


def price_plausible(price, deal):
    if price is None or price <= 0:
        return None
    if (deal or "").lower().startswith("rent"):
        return 3000 <= price <= 500000
    return 800000 <= price <= 300000000


def score_units(data, now):
    pool, seen = [], set()
    skipped = Counter()
    for r in data:
        if str(r.get("availability", "")).lower() != "available":
            skipped["not_available"] += 1
            continue
        phone = clean_phone(r.get("phone"))
        key = f"{phone}_{int(r['price']) if r.get('price') else 0}"
        if key in seen:
            skipped["duplicate"] += 1
            continue
        seen.add(key)
        if not phone:
            skipped["no_phone"] += 1
            continue
        price = r.get("price")
        deal = str(r.get("deal") or "").strip()
        compound = str(r.get("compound") or "").strip() or None
        ptype = str(r.get("propertyType") or "").strip() or None
        beds, baths, area = r.get("rooms"), r.get("bathrooms"), r.get("space")
        updated = parse_date(r.get("lastUpdate"))
        plausible = price_plausible(price, deal)
        score = 25
        score += 25 if plausible else (10 if price else 0)
        score += sum(5 for v in (compound, ptype, beds, baths, area) if v)
        if updated:
            age = (now - updated).days
            score += 15 if age <= 30 else 10 if age <= 90 else 5 if age <= 180 else 0
        score += 10 if compound and compound.lower() in TIER_A_COMPOUNDS else 0
        pool.append({
            "compound": compound, "ptype": ptype or "Unclassified",
            "deal": deal or "Unknown",
            "price": int(price) if price else "",
            "beds": beds or "", "baths": baths or "", "area": area or "",
            "phone": phone, "score": score,
            "tier": "P1" if score >= 70 else "P2" if score >= 55 else "P3",
            "updated_str": updated.strftime("%Y-%m-%d") if updated else "",
            "plausible": plausible,
        })
    pool.sort(key=lambda u: (-u["score"], u["tier"]))
    return pool, skipped


def build_workbook(top, pool_size, skipped, now, top_n):
    wb = Workbook()
    ws = wb.active
    ws.title = f"Top {top_n} Queue"
    headers = ["Rank", "Unit Code", "Compound", "Property Type", "Deal",
               "Price (EGP)", "Beds", "Baths", "Area (sqm)", "Owner Phone",
               "Score", "Tier", "Last Update", "Status", "Call Notes"]
    last_col = len(headers) + 1
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.row_dimensions[1].height = 15
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=last_col)
    ws["B2"] = "Sierra Estates - Owner Verification Campaign Worklist"
    ws["B2"].font = Font(size=16, color=PRIMARY)
    ws["B2"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 32
    ws["B3"] = ("Owner-direct units ranked by publish-readiness. Call in Rank order. "
                "CONFIDENTIAL - contains owner phone numbers, never commit to the repo.")
    ws["B3"].font = Font(size=9, color=NEUTRAL_600)
    header_fill = PatternFill("solid", fgColor=PRIMARY)
    for c, h in enumerate(headers, start=2):
        cell = ws.cell(row=4, column=c, value=h)
        cell.fill = header_fill
        cell.font = Font(size=11, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[4].height = 28
    for i, u in enumerate(top):
        row = 5 + i
        vals = [i + 1, f"SE-WL-{i + 1:04d}", u["compound"] or "Unclassified",
                u["ptype"], u["deal"], u["price"], u["beds"], u["baths"],
                u["area"], u["phone"], u["score"], u["tier"],
                u["updated_str"], "Pending", ""]
        fill = PatternFill("solid", fgColor="FFFFFF" if i % 2 == 0 else NEUTRAL_100)
        for c, v in enumerate(vals, start=2):
            cell = ws.cell(row=row, column=c, value=v)
            cell.fill = fill
            cell.font = Font(size=11, color=NEUTRAL_900)
        ws.cell(row=row, column=7).number_format = "#,##0"
        for c in (7, 8, 9, 10, 12):
            ws.cell(row=row, column=c).alignment = Alignment(horizontal="right", vertical="center")
        ws.row_dimensions[row].height = 22
    n = len(top)
    dv = DataValidation(type="list", formula1=f'"{STATUS_OPTIONS}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"O5:O{4 + n}")
    ws.conditional_formatting.add(
        f"L5:L{4 + n}", DataBarRule(start_type="num", start_value=0,
                                    end_type="num", end_value=100,
                                    color=PRIMARY, showValue=True))
    ws.freeze_panes = "D5"
    for c in range(2, last_col + 1):
        width = max(
            (len(str(ws.cell(row=r, column=c).value or "")) for r in range(5, min(5 + n, 60))),
            default=8)
        ws.column_dimensions[get_column_letter(c)].width = min(max(width + 2, 9), 24)

    ws2 = wb.create_sheet("Call Script")
    ws2.sheet_view.showGridLines = False
    ws2.column_dimensions["A"].width = 3
    script_rows = [
        ("Opening (EN)", "Hello, this is Sierra Estates property management. We list "
         "verified owner-direct properties in New Cairo. Are you still offering your "
         "unit for {deal}?"),
        ("Identity confirm", "Confirm you are the owner and the unit is yours to list."),
        ("Availability", "Is it still available? Expected timeline if rented/sold soon?"),
        ("Price capture", "What is your current asking price? (rent: monthly EGP / sale: "
         "total EGP). Any flexibility?"),
        ("Layout capture", "How many bedrooms and bathrooms? Total area in sqm?"),
        ("Finishing capture", "Finishing level: fully finished / semi / core & shell? Furnished?"),
        ("Photos capture", "Can you send 3+ recent photos on WhatsApp right now?"),
        ("Viewing consent", "Would you accept viewing requests we forward to you?"),
        ("Close", "Thank you - we will list your unit as VERIFIED and route matched "
         "buyers to you directly."),
        ("WhatsApp template", "Hello, this is Sierra Estates. Is your property in {compound} "
         "still available for {deal}? We have verified clients actively looking. Reply YES "
         "and we will handle the rest."),
        ("No-answer policy", "Two attempts on different days, then mark 'Call Again'."),
        ("Opt-out policy", "Mark 'Opted Out' immediately and never re-contact."),
    ]
    ws2["B2"] = "Verification Call Script (English)"
    ws2["B2"].font = Font(size=16, color=PRIMARY)
    for c, h in enumerate(["Section", "Content"], start=2):
        cell = ws2.cell(row=4, column=c, value=h)
        cell.fill = header_fill
        cell.font = Font(size=11, color="FFFFFF")
    for i, (a, b) in enumerate(script_rows):
        row = 5 + i
        fill = PatternFill("solid", fgColor="FFFFFF" if i % 2 == 0 else NEUTRAL_100)
        ca = ws2.cell(row=row, column=2, value=a)
        cb = ws2.cell(row=row, column=3, value=b)
        for cell in (ca, cb):
            cell.fill = fill
            cell.font = Font(size=11, color=NEUTRAL_900)
        cb.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws2.row_dimensions[row].height = 34
    ws2.column_dimensions["B"].width = 20
    ws2.column_dimensions["C"].width = 90

    ws3 = wb.create_sheet("Summary")
    ws3.sheet_view.showGridLines = False
    ws3.column_dimensions["A"].width = 3
    p1 = sum(1 for u in top if u["tier"] == "P1")
    plausible = sum(1 for u in top if u["plausible"])
    rent = sum(1 for u in top if u["deal"].lower().startswith("rent"))
    kpis = [
        ("Source units (owner-direct)", len(pool_size)),
        ("Queue pool (valid phone, deduped, available)", len(pool_size) - sum(
            v for k, v in skipped.items())),
        ("Top queue size", len(top)),
        ("P1 near-publishable (score >= 70)", p1),
        ("Price already plausible in top", plausible),
        ("Rent vs Sale in top", f"{rent} rent / {len(top) - rent} sale"),
        ("Expected calls/day (sierra_master quota)", 30),
        ("Days to clear queue", round(len(top) / 30)),
    ]
    ws3["B2"] = "Campaign Summary"
    ws3["B2"].font = Font(size=16, color=PRIMARY)
    for c, h in enumerate(["Metric", "Value"], start=2):
        cell = ws3.cell(row=4, column=c, value=h)
        cell.fill = header_fill
        cell.font = Font(size=11, color="FFFFFF")
    for i, (k, v) in enumerate(kpis):
        row = 5 + i
        fill = PatternFill("solid", fgColor="FFFFFF" if i % 2 == 0 else NEUTRAL_100)
        ws3.cell(row=row, column=2, value=k).fill = fill
        cv = ws3.cell(row=row, column=3, value=v)
        cv.fill = fill
        cv.alignment = Alignment(horizontal="right", vertical="center")
        for cell in (ws3.cell(row=row, column=2), cv):
            cell.font = Font(size=11, color=NEUTRAL_900)
    compounds = Counter(u["compound"] or "Unclassified" for u in top).most_common(10)
    start = 5 + len(kpis) + 2
    ws3.cell(row=start - 1, column=2, value="Top compounds in queue").font = Font(
        size=12, color=PRIMARY)
    for c, h in enumerate(["Compound", "Units"], start=2):
        cell = ws3.cell(row=start, column=c, value=h)
        cell.fill = header_fill
        cell.font = Font(size=11, color="FFFFFF")
    for i, (name, cnt) in enumerate(compounds):
        row = start + 1 + i
        fill = PatternFill("solid", fgColor="FFFFFF" if i % 2 == 0 else NEUTRAL_100)
        ws3.cell(row=row, column=2, value=name).fill = fill
        cc = ws3.cell(row=row, column=3, value=cnt)
        cc.fill = fill
        cc.alignment = Alignment(horizontal="right", vertical="center")
    chart = BarChart()
    chart.type = "bar"
    chart.title = "Top compounds in queue"
    chart.add_data(Reference(ws3, min_col=3, min_row=start, max_row=start + len(compounds)),
                   titles_from_data=True)
    chart.set_categories(Reference(ws3, min_col=2, min_row=start + 1,
                                   max_row=start + len(compounds)))
    chart.legend = None
    ws3.add_chart(chart, f"E{start}")
    ws3.column_dimensions["B"].width = 44
    ws3.column_dimensions["C"].width = 14
    return wb


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--top", type=int, default=500)
    ap.add_argument("--out-dir", default=os.path.join(REPO, "..", "sierra-campaign"))
    args = ap.parse_args()
    if not os.path.exists(SRC):
        sys.exit(f"source not found: {SRC}")
    data = json.load(open(SRC, encoding="utf-8"))
    now = datetime.now(timezone.utc).replace(tzinfo=None)
    pool, skipped = score_units(data, now)
    top = pool[: args.top]
    os.makedirs(args.out_dir, exist_ok=True)
    xlsx_out = os.path.join(args.out_dir, f"Sierra_Verification_Campaign_Top{args.top}.xlsx")
    csv_out = os.path.join(args.out_dir, f"Sierra_Verification_Campaign_Top{args.top}.csv")
    wb = build_workbook(top, data, skipped, now, args.top)
    wb.properties.creator = "Sierra Estates automation"
    wb.save(xlsx_out)
    with open(csv_out, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Rank", "Unit Code", "Compound", "Property Type", "Deal",
                    "Price (EGP)", "Beds", "Baths", "Area (sqm)", "Owner Phone",
                    "Score", "Tier", "Last Update", "Status", "Call Notes"])
        for i, u in enumerate(top):
            w.writerow([i + 1, f"SE-WL-{i + 1:04d}", u["compound"] or "Unclassified",
                        u["ptype"], u["deal"], u["price"], u["beds"], u["baths"],
                        u["area"], u["phone"], u["score"], u["tier"],
                        u["updated_str"], "Pending", ""])
    print(f"pool={len(pool)} skipped={dict(skipped)}")
    print(f"tiers in top {len(top)}: {Counter(u['tier'] for u in top)}")
    print(f"price-plausible in top: {sum(1 for u in top if u['plausible'])}/{len(top)}")
    print(f"wrote {xlsx_out}")
    print(f"wrote {csv_out}")
    print("NOTE: output contains owner PII - keep outside the repository.")


if __name__ == "__main__":
    main()
